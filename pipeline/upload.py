# pipeline/upload.py
"""上传资料进知识库：解析 → 落源 → 切块 → 嵌入 → 入库。

支持 txt / md / csv / Excel(xlsx,xlsm,xls) / Word(docx)。老格式 .doc/.wps
不读（OLE 二进制，没有可靠的纯 Python 读法），报错提示用户另存为。

**三条设计原则**：

1. **源文件必须留一份，而且留原文件。** 上传的东西复制到
   ``data/chat_raw/uploaded/``，**原样字节复制**（xlsx 还是 xlsx）。
   不存"解析后的文本" —— 那样「重建索引」重新解析时表格结构就没了，
   切块方式跟着变，检索质量静默下降，用户完全不会知道。
   索引是可再生的，源才是资产。

2. **一行一块，块里带表头。** 整张表嵌成一块时，"索尼A7M4一天多少钱"这种
   问句会被几十行无关内容稀释 —— 实测余弦 0.63，够不到 0.65 的直发门槛，
   于是"明明传了价目表却还是转人工"。所以表格按行切，每块前面重复表头。

3. **一个文件失败不影响其它文件。** 每个文件独立 try，失败原因单独返回给界面。
   批量上传里最气的就是"第 3 个文件编码不对，导致 1、2 个也没进去"。
"""

from __future__ import annotations

import csv
import datetime
import io
import logging
import os
import re
import shutil
from pathlib import Path
from typing import Callable, List, Optional, Tuple

logger = logging.getLogger(__name__)

HERE = Path(__file__).resolve().parent.parent

# 源目录跟资料库档案走：第一个档案 data/chat_raw/uploaded/（老位置，不搬家），
# 新建的档案 data/chat_raw/archives/<档案>/。显式赋值可覆盖（测试用）。
UPLOAD_DIR: Optional[Path] = None


def upload_dir() -> Path:
    """当前档案的源文件目录。**每次调用现算**，切档案立刻生效。"""
    if UPLOAD_DIR is not None:
        return Path(UPLOAD_DIR)
    try:
        from rag import archives
        return archives.upload_dir()
    except Exception:
        return HERE / "data" / "chat_raw" / "uploaded"


def active_collection() -> str:
    """当前档案的 collection —— 上传/重建只作用于它，不会清到别的档案。"""
    try:
        from rag import archives
        return archives.active_collection()
    except Exception:
        return "knowledge_base"


# 纯文本 + 表格 + Word。顺序就是界面筛选器的顺序。
SUPPORTED = (".txt", ".md", ".csv", ".xlsx", ".xlsm", ".xls", ".docx")

# 表格格式：按行切块
TABLE_SUFFIXES = (".csv", ".xlsx", ".xlsm", ".xls")

# Word 老格式（.doc）是 OLE 二进制，没有纯 Python 库能可靠读出来。
# 与其装一堆依赖最后读出乱码喂给 LLM，不如直接让用户另存为。
LEGACY_HINT = {
    ".doc": "Word 老格式（.doc）读不出来。用 Word 打开 →「另存为」→ 选 .docx 再上传",
    ".wps": "WPS 老格式读不出来。用 WPS 打开 →「另存为」→ .docx 或 .xlsx 再上传",
}

# 一张表超过这么多行就不收：嵌入是本机 CPU 跑的，几万行会跑到明天早上。
# 明确报错让用户拆文件，比"默默只导了一半"强 —— 半份价目表比没有更危险。
MAX_TABLE_ROWS = 5000


# ── 文本读取 / 单元格整形 ─────────────────────────────────────────────

def _read_text(path: Path) -> str:
    """按常见编码依次试。中文 Windows 上 GBK 的文本很常见，不能只试 utf-8。"""
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "gbk", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ValueError("编码认不出（试过 utf-8 / gbk），请另存为 UTF-8")


def _cell_str(v) -> str:
    """单元格变字符串。

    表格里的数最容易变成 ``120.0`` 或 ``2024-01-01 00:00:00``，这两种写法
    喂给 LLM 都会让它把价格和日期说错，所以在进库前就整形好。
    """
    if v is None:
        return ""
    if isinstance(v, bool):            # bool 是 int 的子类，必须排在数字前面
        return "是" if v else "否"
    if isinstance(v, datetime.datetime):
        return (v.strftime("%Y-%m-%d %H:%M") if (v.hour or v.minute)
                else v.strftime("%Y-%m-%d"))
    if isinstance(v, datetime.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else f"{v:g}"
    return str(v).strip()


def _clean_rows(rows) -> List[Tuple[List[str], List[int]]]:
    """丢掉空行和空单元格，返回 ``[(非空单元格, 原列号), ...]``。

    **为什么记列号**：老板的表里"备注"经常是空的，那一行就少一列。如果只看
    位置对齐，整行标签会全错位甚至全丢（`佳能R5，150，3000` 看不出 150 是日租
    还是押金）。按原列号对表头，缺哪列就少哪列，剩下的照旧对得上。
    """
    out = []
    for row in rows:
        cells, idxs = [], []
        for i, c in enumerate(row):
            s = _cell_str(c)
            if s:
                cells.append(s)
                idxs.append(i)
        if cells:
            out.append((cells, idxs))
    return out


def _join(rows) -> str:
    return "\n".join(" | ".join(cells) for cells, _ in rows)


# 表格行里的数字要认得出来（"120" 是数字，"备注" 不是）
_DIGIT = re.compile(r"\d")

# 「日租/周租/月租」这种表头，客户嘴上说的是「一天/一周/一个月」。
# 不加这层同义词，问「一天多少钱」会被**周租**那块抢到第一 —— 实测：
#   【日租】…日租 120     vs 问「索尼A7M4一天多少钱」→ 输给 【周租】…周租 600
#   【日租】…日租（一天）120 vs 同一问           → 赢，且「一周」问句也归周租
# 只是把表头换个说法，**不加任何数据里没有的价格/单位**。
_UNIT_SYNONYM = (("日租", "一天"), ("周租", "一周"), ("月租", "一个月"),
                 ("季租", "一个季度"), ("年租", "一年"))


def _synonym(header: str) -> str:
    for k, s in _UNIT_SYNONYM:
        if k in header:
            return s
    return ""


def _looks_like_header(cells: List[str]) -> bool:
    """第一行是不是表头。表头都很短、且不含数字。"""
    return (len(cells) >= 2
            and all(len(c) <= 8 and not _DIGIT.search(c) for c in cells))


def _row_sentence(cells: List[str], idxs: List[int], headers: Optional[dict] = None) -> str:
    """一行 → 一句人话：``索尼A7M4：日租（一天）120，押金 2000，含24-70``。

    **为什么不是 ``" | "`` 拼**：问句和竖线表格的余弦明显偏低。实测同一行
    （日租 120 / 押金 2000）对「索尼A7M4一天多少钱」：

    ======================  ======
    写法                    余弦
    ======================  ======
    ``a | 120 | 2000``      0.68
    ``a：日租 120，…``       0.75
    ======================  ======

    0.68 就在 0.65 门槛附近晃，稍微换个问法就掉下去 → 明明有价目表还是转人工。
    """
    if not headers or len(cells) == 1:
        return "，".join(cells)
    parts = []
    for c, i in zip(cells[1:], idxs[1:]):
        h = headers.get(i, "")
        if not h:
            parts.append(c)
            continue
        syn = _synonym(h)
        parts.append(f"{h}（{syn}）{c}" if syn else f"{h} {c}")
    return cells[0] + ("：" + "，".join(parts) if parts else "")


# ── 表格 → 可读文本 + 行级分块 ────────────────────────────────────────

def _table_from_rows(rows, label: str = "") -> Tuple[List[str], List[str]]:
    """表格 → ``(可读文本行, 分块)``。第一行当表头，之后每行一块、每块带表名。

    每块都带 ``label`` 是因为**日租和周租会打架**：A7M4 日租 120、周租 600，
    不写清是哪张表，机器人可能把 600 当成一天的价格报出去。
    """
    if not rows:
        return [], []
    lines: List[str] = []
    if label:
        lines.append(f"【{label}】")
    lines += [" | ".join(cells) for cells, _ in rows]
    prefix = f"【{label}】" if label else ""

    headers = dict(zip(rows[0][1], rows[0][0])) if _looks_like_header(rows[0][0]) else None
    data = rows[1:] if headers else rows
    if not data:                      # 光有表头没有数据行
        return lines, [prefix + " | ".join(rows[0][0])]
    return lines, [prefix + _row_sentence(c, i, headers) for c, i in data]


def _guard_rows(chunks: List[str]) -> List[str]:
    if len(chunks) > MAX_TABLE_ROWS:
        raise ValueError(f"表格太大（{len(chunks)} 行，上限 {MAX_TABLE_ROWS} 行）。"
                         f"请拆成多个文件再传，不然嵌入要跑很久")
    return chunks


# ── 各格式解析：都返回 (可读文本, 分块或 None) ────────────────────────

def _parse_xlsx(p: Path) -> Tuple[str, List[str]]:
    """Excel 2007+（.xlsx / .xlsm）。read_only 模式省内存，大表也不会撑爆。"""
    import openpyxl

    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    try:
        lines: List[str] = []
        chunks: List[str] = []
        multi = len(wb.sheetnames) > 1
        for name in wb.sheetnames:
            rows = _clean_rows(wb[name].iter_rows(values_only=True))
            if not rows:
                continue                    # 空工作表直接跳过，别留个光秃秃的表名
            l, c = _table_from_rows(rows, name if multi else "")
            lines += l
            chunks += c
            _guard_rows(chunks)
        return "\n".join(lines), chunks
    finally:
        wb.close()


def _parse_xls(p: Path) -> Tuple[str, List[str]]:
    """Excel 2003（.xls）。xlrd 只管这一个格式。"""
    import xlrd

    book = xlrd.open_workbook(str(p))
    lines: List[str] = []
    chunks: List[str] = []
    multi = book.nsheets > 1
    for sh in book.sheets():
        raw = []
        for r in range(sh.nrows):
            cells = []
            for c in range(sh.ncols):
                v = sh.cell_value(r, c)
                if sh.cell_type(r, c) == xlrd.XL_CELL_DATE:
                    try:
                        v = xlrd.xldate_as_datetime(v, book.datemode)
                    except Exception:
                        pass            # 认不出的日期当普通值，别让整个文件失败
                cells.append(_cell_str(v))
            raw.append(cells)
        rows = _clean_rows(raw)
        if not rows:
            continue
        l, c = _table_from_rows(rows, sh.name if multi else "")
        lines += l
        chunks += c
        _guard_rows(chunks)
    return "\n".join(lines), chunks


def _parse_docx(p: Path) -> Tuple[str, List[str]]:
    """Word（.docx）。段落按语义切，表格按行切 —— 价目表几乎都放在表格里。"""
    import docx

    from pipeline.chunker import split_text

    d = docx.Document(str(p))
    lines: List[str] = []
    chunks: List[str] = []

    paras = [t for t in (para.text.strip() for para in d.paragraphs) if t]
    if paras:
        lines += paras
        chunks += [c for c in split_text("\n".join(paras)) if c.strip()]

    tables = d.tables
    for i, tb in enumerate(tables, 1):
        rows = _clean_rows([[cell.text for cell in row.cells] for row in tb.rows])
        if not rows:
            continue
        l, c = _table_from_rows(rows, f"表格{i}" if len(tables) > 1 else "表格")
        lines += l
        chunks += c
        _guard_rows(chunks)
    return "\n".join(lines), chunks


def parse_records(path: os.PathLike) -> Tuple[str, Optional[List[str]]]:
    """解析文件 → ``(可读文本, 分块)``。``分块=None`` 表示交给通用切块器。

    分块单独返回是因为表格不能整段切：整张表一块会稀释检索（见模块头原则 2）。
    """
    p = Path(path)
    if not p.is_file():
        raise ValueError("文件不存在")
    suffix = p.suffix.lower()
    if suffix in LEGACY_HINT:
        raise ValueError(LEGACY_HINT[suffix])
    if suffix not in SUPPORTED:
        raise ValueError(f"暂不支持 {suffix or '无扩展名'}（支持 "
                         f"{' / '.join(SUPPORTED)}）")

    if suffix == ".csv":
        rows = _clean_rows(csv.reader(io.StringIO(_read_text(p))))
        lines, chunks = _table_from_rows(rows)
        return "\n".join(lines), _guard_rows(chunks)
    if suffix in (".xlsx", ".xlsm"):
        return _parse_xlsx(p)
    if suffix == ".xls":
        return _parse_xls(p)
    if suffix == ".docx":
        return _parse_docx(p)

    return _read_text(p), None          # txt / md：交给通用切块器


def parse_file(path: os.PathLike) -> str:
    """只要文本（预览 / 测试用）。不支持的格式抛 ValueError。"""
    return parse_records(path)[0]


# ── 留源（原文件原样复制） ────────────────────────────────────────────

def save_source(path: os.PathLike) -> Path:
    """原文件复制一份到 uploaded/，重名加 ``_1``。"""
    d = upload_dir()
    d.mkdir(parents=True, exist_ok=True)
    p = Path(path)
    dest = d / p.name
    n = 1
    while dest.exists():
        dest = d / f"{p.stem}_{n}{p.suffix}"
        n += 1
    shutil.copy2(p, dest)
    return dest


def backup_dir() -> Path:
    """被替换掉的旧源挪这里，不删 —— 换错了还能捞回来。"""
    return upload_dir() / ".old"


def replace_source(path: os.PathLike) -> Tuple[Path, Optional[Path]]:
    """留源，并处理"同一个文件重传"。返回 ``(新源, 被顶掉的旧源或 None)``。

    **为什么必须处理重传**：老板改了价目表再传一次，如果只往后加，库里就同时有
    新旧两份价格，检索可能命中旧的 —— 机器人报错价，比不回答严重得多。
    所以同名源＝替换：旧源挪到 ``.old/``，旧向量由调用方删掉。
    """
    d = upload_dir()
    d.mkdir(parents=True, exist_ok=True)
    p = Path(path)
    dest = d / p.name
    old = dest if dest.exists() else None
    if old is not None:
        b = backup_dir()
        b.mkdir(parents=True, exist_ok=True)
        stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        try:
            old.replace(b / f"{p.stem}_{stamp}{p.suffix}")
        except OSError as e:                      # 挪不动就直接覆盖，别让上传失败
            logger.warning(f"旧源备份失败（{old}）: {e}")
    shutil.copy2(p, dest)
    return dest, old


def already_home(p: Path) -> bool:
    """这个文件本来就在源目录里吗（重建索引走的就是这条路）。

    重建时**绝不能**再复制一遍、更不能当"重传"把源挪进 ``.old/`` ——
    那等于一边重建一边把源删了。
    """
    try:
        rp = p.resolve()
    except OSError:
        return False
    raw = (HERE / "data" / "chat_raw").resolve()
    return rp.parent in (upload_dir().resolve(), raw)


def list_uploaded() -> List[dict]:
    """已上传的源文件（界面展示用）。"""
    d = upload_dir()
    if not d.is_dir():
        return []
    out = []
    for f in sorted(d.iterdir()):
        if not f.is_file() or f.suffix.lower() not in SUPPORTED:
            continue                       # 跳过 .old 目录之类
        st = f.stat()
        out.append({"name": f.name,
                    "size_kb": round(st.st_size / 1024, 1),
                    "mtime": datetime.datetime.fromtimestamp(st.st_mtime)
                                     .strftime("%Y-%m-%d %H:%M")})
    return out


# ── 删旧向量 ──────────────────────────────────────────────────────────

def delete_source_points(qdrant, collection: str, source: str) -> int:
    """删掉某个源产生的向量，返回删掉的条数；失败返回 ``-1``。

    按 ``payload.source`` 删 —— 这也是 embedder 写 payload 时带上 source 的用处。
    """
    try:
        from qdrant_client.models import (FieldCondition, Filter, FilterSelector,
                                          MatchValue)
        flt = Filter(must=[FieldCondition(key="source", match=MatchValue(value=source))])
        n = qdrant.count(collection_name=collection, count_filter=flt, exact=True).count
        qdrant.delete(collection_name=collection,
                      points_selector=FilterSelector(filter=flt), wait=True)
        return n
    except Exception as e:
        logger.warning(f"删旧向量失败（source={source}）: {e}")
        return -1


# ── 批量导入 ──────────────────────────────────────────────────────────

def import_files(paths: List[os.PathLike], qdrant, collection: str = None,
                 progress: Optional[Callable[[str], None]] = None) -> dict:
    """批量导入。返回 ``{"ok": [...], "failed": [{file, reason}], "chunks": n}``。

    ``progress`` 每步回调一次（界面用来刷状态），**跑在后台线程里**。
    同名文件重传＝替换（旧源备份、旧向量删除），不会新旧价格并存。
    """
    from pipeline.chunker import split_text
    from pipeline.embedder import embed_and_store
    collection = collection or active_collection()

    def say(msg):
        logger.info(msg)
        if progress:
            try:
                progress(msg)
            except Exception:
                pass

    ok, failed, total = [], [], 0
    seen_stems = set()
    for path in paths:
        p = Path(path)
        try:
            text, chunks = parse_records(p)
            if not text.strip():
                raise ValueError("文件里没有可读文字")
            if already_home(p):
                src, old = p, None          # 重建索引：源就在源目录里，别再复制
            elif p.stem in seen_stems:
                # 同一批里两个不同目录的同名文件：不能互相顶掉，加后缀共存
                src, old = save_source(p), None
            else:
                seen_stems.add(p.stem)
                src, old = replace_source(p)

            replaced = 0
            if old is not None:
                replaced = delete_source_points(qdrant, collection, old.name)
                if replaced < 0:
                    say(f"{p.name}: 旧版本索引没删干净，建议点「重建索引」")
                else:
                    say(f"{p.name}: 换了新版本，旧索引清掉 {replaced} 条")

            if chunks is None:
                chunks = [c for c in split_text(text) if c.strip()]
            if not chunks:
                raise ValueError("切不出有效内容")
            say(f"{p.name}: 切出 {len(chunks)} 块，正在嵌入…")
            n = embed_and_store(chunks, qdrant, collection_name=collection,
                                source=src.name)
            total += n
            ok.append({"file": p.name, "chunks": n, "source": src.name,
                       "replaced": old.name if old is not None else None,
                       "removed_points": max(0, replaced)})
            say(f"{p.name}: 已入库 {n} 块")
        except Exception as e:
            failed.append({"file": p.name, "reason": f"{type(e).__name__}: {e}"})
            say(f"{p.name}: 失败 —— {e}")
    return {"ok": ok, "failed": failed, "chunks": total}


def source_files() -> List[Path]:
    """**当前档案**可重建索引的源文件。

    ``data/chat_raw/*.txt`` 是最早那条导入管线留下的源（视频转录/FAQ），
    它们属于**第一个档案**。别的档案绝不能带进来 —— 否则在"服装租赁"点重建索引，
    会把摄影器材的资料灌进去（串味 = 报错价）。
    """
    out: List[Path] = []
    if is_first_archive():
        raw = HERE / "data" / "chat_raw"
        if raw.is_dir():
            out += [f for f in sorted(raw.glob("*.txt")) if f.is_file()]
    ud = upload_dir()
    if ud.is_dir():
        out += [f for f in sorted(ud.iterdir())
                if f.is_file() and f.suffix.lower() in SUPPORTED]
    return out


def is_first_archive() -> bool:
    """当前档案是不是从老库升上来的那一个。"""
    try:
        from rag import archives
        return archives.is_first(archives.active_id())
    except Exception:
        return True
